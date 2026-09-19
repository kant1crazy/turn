"""Lecture/écriture de data/articles.xlsx + calcul du prix conseillé.

Règle de prix (décidée le 19/09/2026) :
- toujours disponible, sans risque : prix_achat × marge par défaut ;
- si la case "utiliser_prix_marche" de la ligne est cochée (Oui) : on
  mélange avec la médiane des prix constatés sur Vinted pour la même
  marque/type/couleur/taille (voir scraper.py, qui peut échouer — dans ce
  cas on retombe silencieusement sur le mode par défaut, annoncé dans le
  champ `detail`).
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from . import config, scraper
from .models import ArticleEtat, ArticleInfo, ArticleType, PriceSuggestion

# (clé, en-tête) — source unique utilisée par le script de génération du
# template et par la lecture/écriture des lignes.
COLUMNS = [
    ("id", "ID"),
    ("date_ajout", "Date d'ajout"),
    ("type", "Type"),
    ("marque", "Marque"),
    ("couleur", "Couleur"),
    ("taille", "Taille"),
    ("etat", "État"),
    ("matiere", "Matière"),
    ("prix_achat", "Prix d'achat (€)"),
    ("utiliser_prix_marche", "Utiliser le prix marché"),
    ("prix_marche_constate", "Prix marché constaté (€)"),
    ("nb_annonces_comparables", "Nb annonces comparables"),
    ("prix_conseille", "Prix conseillé (€)"),
    ("prix_vente_final", "Prix de vente final (€)"),
    ("notes", "Notes"),
]
COL_INDEX = {key: i + 1 for i, (key, _) in enumerate(COLUMNS)}


def build_template() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "articles"
    for i, (_, header) in enumerate(COLUMNS, start=1):
        ws.cell(row=1, column=i, value=header)
    ws.freeze_panes = "A2"

    type_dv = DataValidation(
        type="list", formula1='"' + ",".join(t.value for t in ArticleType) + '"', allow_blank=True
    )
    etat_dv = DataValidation(
        type="list", formula1='"' + ",".join(e.value for e in ArticleEtat) + '"', allow_blank=True
    )
    bool_dv = DataValidation(type="list", formula1='"Oui,Non"', allow_blank=True)
    for dv in (type_dv, etat_dv, bool_dv):
        ws.add_data_validation(dv)
    type_dv.add(f"{_col_letter('type')}2:{_col_letter('type')}1000")
    etat_dv.add(f"{_col_letter('etat')}2:{_col_letter('etat')}1000")
    bool_dv.add(f"{_col_letter('utiliser_prix_marche')}2:{_col_letter('utiliser_prix_marche')}1000")

    widths = {"marque": 16, "couleur": 12, "notes": 30, "date_ajout": 14}
    for key, _ in COLUMNS:
        letter = _col_letter(key)
        ws.column_dimensions[letter].width = widths.get(key, 18)
    return wb


def _col_letter(key: str) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(COL_INDEX[key])


def ensure_workbook() -> Path:
    if not config.EXCEL_PATH.exists():
        config.EXCEL_PATH.parent.mkdir(parents=True, exist_ok=True)
        build_template().save(config.EXCEL_PATH)
    return config.EXCEL_PATH


def _find_row(ws, article_id: str) -> int | None:
    id_col = COL_INDEX["id"]
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=id_col).value == article_id:
            return row
    return None


def upsert_article_row(article: ArticleInfo, suggestion: PriceSuggestion) -> None:
    ensure_workbook()
    wb = load_workbook(config.EXCEL_PATH)
    ws = wb["articles"]
    row = _find_row(ws, article.id) or (ws.max_row + 1 if ws.cell(row=2, column=1).value else 2)

    values = {
        "id": article.id,
        "date_ajout": date.today().isoformat(),
        "type": article.type.value,
        "marque": article.marque,
        "couleur": article.couleur,
        "taille": article.taille,
        "etat": article.etat.value,
        "matiere": article.matiere or "",
        "prix_achat": article.prix_achat,
        "utiliser_prix_marche": "Oui" if article.utiliser_prix_marche else "Non",
        "prix_marche_constate": suggestion.prix_marche_constate or "",
        "nb_annonces_comparables": suggestion.nb_annonces_comparables,
        "prix_conseille": suggestion.prix_conseille,
    }
    for key, value in values.items():
        ws.cell(row=row, column=COL_INDEX[key], value=value)
    wb.save(config.EXCEL_PATH)


def _round_to_vinted(x: float) -> float:
    return round(x * 2) / 2  # arrondi au 0,50 € le plus proche


def compute_suggested_price(article: ArticleInfo) -> PriceSuggestion:
    base = article.prix_achat * config.DEFAULT_MARGIN_MULTIPLIER
    detail = f"Prix d'achat × {config.DEFAULT_MARGIN_MULTIPLIER} = {round(base, 2)} €"
    market_price = None
    n = 0

    if article.utiliser_prix_marche:
        try:
            prices = scraper.search_market_prices(
                marque=article.marque, couleur=article.couleur, taille=article.taille, type_hint=article.type.value
            )
            market_price = scraper.median_price(prices)
            n = len(prices)
        except scraper.ScraperError as exc:
            detail += f" — estimation marché indisponible ({exc}), prix basé sur la marge par défaut"

    if market_price:
        suggested = max(article.prix_achat, (base + market_price) / 2)
        detail = (
            f"Moyenne entre marge par défaut ({round(base, 2)} €) et médiane "
            f"marché sur {n} annonces comparables ({market_price} €)"
        )
    else:
        suggested = base

    suggested = _round_to_vinted(suggested)
    return PriceSuggestion(
        prix_achat=article.prix_achat,
        utiliser_prix_marche=article.utiliser_prix_marche,
        prix_marche_constate=market_price,
        nb_annonces_comparables=n,
        prix_conseille=suggested,
        detail=detail,
    )
